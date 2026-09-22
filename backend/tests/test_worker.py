from docx import Document

from app.extraction import parse_document, suggest_evidence
from app.worker import valid_file_signature


def test_file_signatures_match_declared_type() -> None:
    assert valid_file_signature("application/pdf", b"%PDF")
    assert valid_file_signature(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        b"PK\x03\x04",
    )
    assert not valid_file_signature("application/pdf", b"PK\x03\x04")
    assert not valid_file_signature("application/octet-stream", b"%PDF")


def test_docx_parsing_and_evidence_suggestions(tmp_path) -> None:
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_heading("Experience", level=1)
    document.add_paragraph("Built a FastAPI service with PostgreSQL and Docker, deployed through GitHub Actions.")
    document.save(path)

    parsed = parse_document(
        path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    suggestions = {item.skill: item for item in suggest_evidence(parsed.text)}

    assert parsed.character_count > 20
    assert len(parsed.text_sha256) == 64
    assert suggestions["FastAPI"].proposed_level == 3
    assert suggestions["PostgreSQL"].excerpt.startswith("Built")
    assert suggestions["Docker"].confidence == 0.9
    assert "GitHub Actions" in suggestions
