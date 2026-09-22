import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .taxonomy import SKILLS

PARSER_VERSION = "local-taxonomy-v1"

@dataclass(frozen=True)
class ParsedDocument:
    text: str
    text_sha256: str
    character_count: int
    page_count: int | None


@dataclass(frozen=True)
class SuggestedEvidence:
    skill: str
    category: str
    excerpt: str
    locator: str
    confidence: float
    proposed_level: int


def parse_document(path: Path, content_type: str) -> ParsedDocument:
    if content_type == "application/pdf":
        reader = PdfReader(path)
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n".join(page for page in pages if page)
        page_count = len(reader.pages)
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        document = Document(path)
        text = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
        page_count = None
    else:
        raise ValueError("Unsupported document type")
    normalized = re.sub(r"[ \t]+", " ", text).strip()
    if len(normalized) < 20:
        raise ValueError("Document contains too little extractable text")
    return ParsedDocument(
        text=normalized,
        text_sha256=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        character_count=len(normalized),
        page_count=page_count,
    )


def suggest_evidence(text: str) -> list[SuggestedEvidence]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    suggestions: list[SuggestedEvidence] = []
    for _, (skill, category, patterns) in SKILLS.items():
        match_line = next((line for line in lines if any(re.search(pattern, line, re.I) for pattern in patterns)), None)
        if match_line is None:
            continue
        excerpt = match_line[:400]
        action_pattern = r"\b(built|developed|implemented|deployed|designed|created|automated|integrated)\b"
        implementation_signal = bool(re.search(action_pattern, excerpt, re.I))
        suggestions.append(
            SuggestedEvidence(
                skill=skill,
                category=category,
                excerpt=excerpt,
                locator=f"line {lines.index(match_line) + 1}",
                confidence=0.9 if implementation_signal else 0.75,
                proposed_level=3 if implementation_signal else 2,
            )
        )
    return suggestions
